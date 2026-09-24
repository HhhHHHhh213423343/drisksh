from __future__ import annotations

import asyncio
import base64
import hashlib
import sys
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session


ROOT = Path(__file__).resolve().parents[1]
BACKEND_DIR = ROOT / "backend"
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

from app.db.base import Base
from app.config import get_settings
from app.api.routes.company_profile import (
    export_company_profile,
    require_collection_key,
)
from app.models import Company, RelatedEntity, RiskEvent
from app.schemas.company import CompanySearchIngestRequest
from app.schemas.company_profile import CompanyProfileComplete
from app.services.company_profile import (
    PROFILE_MODULE_KEYS,
    claim_next_run,
    complete_run,
    create_or_reuse_run,
    heartbeat_run,
    profile_company_suggestion,
    resolve_profile_company_name,
)
from app.services.demo_scope import DemoScopeError, canonical_company_name
from enterprise_sentinel.company_profile.collector import (
    _business_change_needs_review,
    sections_from_snapshot,
)
from enterprise_sentinel.company_profile.agent import CompanyProfileAgent, WorkerApiError
from enterprise_sentinel.company_profile.contract import MODULE_BY_KEY, MODULE_SPECS, SOP_VERSION
from enterprise_sentinel.company_profile.excel import build_company_profile_workbook
from enterprise_sentinel.engine import BrowserConfig, CrawlerEngine


def database() -> Session:
    engine = create_engine("sqlite+pysqlite:///:memory:", future=True)
    Base.metadata.create_all(engine)
    return Session(engine)


def sample_normalized(company_name: str) -> dict:
    modules = {}
    for spec in MODULE_SPECS:
        modules[spec.key] = {
            "key": spec.key,
            "title": spec.title,
            "status": "PASS",
            "source_url": spec.url("test-code"),
            "captured_at": datetime.now(timezone.utc).isoformat(),
            "page_count": 1,
            "row_count": 1,
            "sections": [
                {
                    "title": spec.title,
                    "columns": list(spec.empty_columns),
                    "rows": [[str(index + 1) for index in range(len(spec.empty_columns))]],
                }
            ],
        }
    return {
        "company_name": company_name,
        "company_code": "test-code",
        "captured_at": datetime.now(timezone.utc).isoformat(),
        "modules": modules,
    }


def test_run_claim_complete_and_24_hour_reuse() -> None:
    db = database()
    company = Company(name="上海携程金融信息服务有限公司", company_profile={})
    db.add(company)
    db.commit()
    db.refresh(company)

    started = create_or_reuse_run(db, company)
    assert started["enabled"] is True
    assert started["reused"] is False

    duplicate = create_or_reuse_run(db, company)
    assert duplicate["reused"] is True
    assert duplicate["run"]["id"] == started["run"]["id"]

    job = claim_next_run(db, "mac-test")
    assert job and job["company"]["name"] == company.name
    assert job["company"]["qyyjt_profile"]["company_code"] == (
        "C71376C50868F97142E2EDD20C553765"
    )
    run_id = job["run"]["id"]
    heartbeat = heartbeat_run(
        db,
        run_id,
        "mac-test",
        progress_current=3,
        current_module="business_changes",
        module_statuses={"overview": "PASS"},
    )
    assert heartbeat["progress_current"] == 3

    captured_at = datetime.now(timezone.utc)
    normalized = sample_normalized(company.name)
    filename, excel_bytes = build_company_profile_workbook(company.name, normalized, captured_at)
    statuses = {key: "PASS" for key in PROFILE_MODULE_KEYS}
    completed = complete_run(
        db,
        run_id,
        CompanyProfileComplete(
            worker_id="mac-test",
            captured_at=captured_at,
            sop_version=SOP_VERSION,
            company_code="test-code",
            module_statuses=statuses,
            normalized_data=normalized,
            raw_data={"modules": {}},
            excel_filename=filename,
            excel_sha256=hashlib.sha256(excel_bytes).hexdigest(),
            excel_base64=base64.b64encode(excel_bytes).decode("ascii"),
        ),
    )
    assert completed["excel_filename"] == filename

    download = export_company_profile(company.id, db)
    assert download.body == excel_bytes
    assert download.headers["content-disposition"].startswith(
        "attachment; filename*=UTF-8''"
    )

    cached = create_or_reuse_run(db, company)
    assert cached["reused"] is True
    assert cached["run"] is None
    assert cached["snapshot"]["id"] == completed["id"]


def test_incomplete_run_cannot_replace_snapshot() -> None:
    db = database()
    company = Company(name="上海携程金融信息服务有限公司", company_profile={})
    db.add(company)
    db.commit()
    db.refresh(company)
    run = create_or_reuse_run(db, company)["run"]
    claim_next_run(db, "mac-test")
    normalized = sample_normalized(company.name)
    filename, excel_bytes = build_company_profile_workbook(
        company.name,
        normalized,
        datetime.now(timezone.utc),
    )

    with pytest.raises(ValueError, match="八模块尚未全部通过"):
        complete_run(
            db,
            run["id"],
            CompanyProfileComplete(
                worker_id="mac-test",
                captured_at=datetime.now(timezone.utc),
                sop_version=SOP_VERSION,
                module_statuses={"overview": "PASS"},
                normalized_data=normalized,
                excel_filename=filename,
                excel_sha256=hashlib.sha256(excel_bytes).hexdigest(),
                excel_base64=base64.b64encode(excel_bytes).decode("ascii"),
            ),
        )


def test_completed_snapshot_is_idempotently_converted_to_monitoring_evidence() -> None:
    db = database()
    company = Company(name="上海携程金融信息服务有限公司", company_profile={})
    db.add(company)
    db.commit()
    db.refresh(company)

    normalized = sample_normalized(company.name)
    normalized["company_code"] = "C71376C50868F97142E2EDD20C553765"
    normalized["modules"]["dynamic_monitor"]["sections"] = [
        {
            "title": "动态监测",
            "columns": ["日期", "标题", "分类", "重要性", "正负面", "来源"],
            "rows": [["2026-09-20", "新增诉讼事项", "司法风险", "重要", "负面", "法院公告"]],
        }
    ]
    normalized["modules"]["investments"]["sections"] = [
        {
            "title": "对外投资企业",
            "columns": ["企业名称", "投资比例", "企业状态", "行业"],
            "rows": [["上海携程小额贷款有限公司", "100%", "存续", "金融"]],
        }
    ]

    def complete_once(worker: str) -> None:
        run = create_or_reuse_run(db, company, force=True)["run"]
        assert claim_next_run(db, worker)
        captured_at = datetime(2026, 9, 24, tzinfo=timezone.utc)
        filename, excel_bytes = build_company_profile_workbook(company.name, normalized, captured_at)
        complete_run(
            db,
            run["id"],
            CompanyProfileComplete(
                worker_id=worker,
                captured_at=captured_at,
                sop_version=SOP_VERSION,
                company_code=normalized["company_code"],
                module_statuses={key: "PASS" for key in PROFILE_MODULE_KEYS},
                normalized_data=normalized,
                raw_data={},
                excel_filename=filename,
                excel_sha256=hashlib.sha256(excel_bytes).hexdigest(),
                excel_base64=base64.b64encode(excel_bytes).decode("ascii"),
            ),
        )

    complete_once("mac-one")
    first_event_count = db.query(RiskEvent).filter(RiskEvent.company_id == company.id).count()
    first_related_count = db.query(RelatedEntity).filter(RelatedEntity.company_id == company.id).count()
    assert first_event_count >= 1
    assert first_related_count >= 1
    event = db.query(RiskEvent).filter(RiskEvent.title == "新增诉讼事项").one()
    assert event.category == "legal"
    assert event.extra_payload["profile_snapshot_id"]

    complete_once("mac-two")
    assert db.query(RiskEvent).filter(RiskEvent.company_id == company.id).count() == first_event_count
    assert db.query(RelatedEntity).filter(RelatedEntity.company_id == company.id).count() == first_related_count


def test_expired_worker_lease_can_be_reclaimed() -> None:
    db = database()
    company = Company(name="上海携程金融信息服务有限公司", company_profile={})
    db.add(company)
    db.commit()
    db.refresh(company)
    run = create_or_reuse_run(db, company)["run"]
    assert claim_next_run(db, "mac-first") is not None

    from app.models import CompanyProfileRun

    stored = db.get(CompanyProfileRun, run["id"])
    stored.lease_expires_at = datetime(2020, 1, 1, tzinfo=timezone.utc)
    db.commit()

    reclaimed = claim_next_run(db, "mac-second")
    assert reclaimed is not None
    assert reclaimed["run"]["id"] == run["id"]
    assert reclaimed["run"]["worker_id"] == "mac-second"


def test_collection_worker_requires_shared_secret(monkeypatch) -> None:
    settings = get_settings()
    monkeypatch.setattr(settings, "collection_api_key", "test-collection-secret", raising=False)

    with pytest.raises(HTTPException) as exc_info:
        require_collection_key(None)
    assert exc_info.value.status_code == 401

    with pytest.raises(HTTPException) as exc_info:
        require_collection_key("wrong-secret")
    assert exc_info.value.status_code == 401

    assert require_collection_key("test-collection-secret") is None


def test_demo_scope_allows_only_ctrip_finance(monkeypatch) -> None:
    monkeypatch.setattr(
        get_settings(),
        "demo_company_name",
        "上海携程金融信息服务有限公司",
        raising=False,
    )
    assert canonical_company_name("携程金融信息服务有限公司") == "上海携程金融信息服务有限公司"
    with pytest.raises(DemoScopeError):
        canonical_company_name("其他企业有限公司")


def test_excel_contract_contains_exact_eight_sheets() -> None:
    company_name = "上海携程金融信息服务有限公司"
    filename, data = build_company_profile_workbook(
        company_name,
        sample_normalized(company_name),
        datetime(2026, 9, 20, tzinfo=timezone.utc),
    )
    workbook = load_workbook(BytesIO(data), data_only=False)

    assert filename.startswith(f"{company_name}-企业全景信息-")
    assert workbook.sheetnames == [spec.title for spec in MODULE_SPECS]
    for sheet in workbook.worksheets:
        assert sheet["A1"].value == sheet.title
        assert "来源：企业预警通" in str(sheet["A3"].value)
        assert sheet["A1"].fill.fgColor.rgb == "FF009A5A"
        assert sheet["A2"].font.bold is True
        assert sheet["A2"].font.color.rgb == "FF166A45"
        assert sheet["A3"].font.sz == 9
        assert sheet.sheet_view.showGridLines is False
        assert sheet.freeze_panes == {
            "企业速览": "A39",
            "监管处罚": "A4",
            "动态监测": "A7",
            "工商变更": "A7",
            "高管信息": "A7",
            "股东信息": "A10",
            "对外投资企业": "A7",
            "控股子公司": "A7",
        }[sheet.title]
        assert all(
            not (isinstance(cell.value, str) and cell.value.startswith("="))
            for row in sheet.iter_rows()
            for cell in row
        )


def test_split_header_and_body_tables_are_joined() -> None:
    spec = MODULE_BY_KEY["shareholders"]
    snapshot = {
        "tables": [
            {"title": "股东信息", "rows": [["股东名称", "持股比例"]]},
            {"title": "", "rows": [["携程旅游网络技术（上海）有限公司", "100%"]]},
        ]
    }

    sections = sections_from_snapshot(spec, snapshot)

    assert sections[0]["columns"] == ["股东名称", "持股比例"]
    assert sections[0]["rows"] == [["携程旅游网络技术（上海）有限公司", "100%"]]


def test_new_columns_are_preserved_as_right_side_extensions() -> None:
    spec = MODULE_BY_KEY["shareholders"]
    snapshot = {
        "tables": [
            {
                "title": "股东信息",
                "rows": [
                    ["股东名称", "持股比例"],
                    ["携程旅游网络技术（上海）有限公司", "100%", "新增值"],
                ],
            }
        ]
    }

    sections = sections_from_snapshot(spec, snapshot)

    assert sections[0]["columns"] == ["股东名称", "持股比例", "扩展字段1"]
    assert sections[0]["rows"][0][-1] == "新增值"


def test_shareholder_controller_notes_are_preserved_before_table() -> None:
    spec = MODULE_BY_KEY["shareholders"]
    snapshot = {
        "keyValues": [
            ["疑似实际控制人", "梁建章"],
            ["控制路径", "梁建章 → 携程集团 → 上海携程金融"],
        ],
        "tables": [
            {
                "title": "股东明细",
                "rows": [
                    ["股东名称", "持股比例"],
                    ["携程旅游网络技术（上海）有限公司", "58.8%"],
                ],
            }
        ],
    }

    sections = sections_from_snapshot(spec, snapshot)

    assert [section.get("kind") for section in sections[:2]] == ["note", "note"]
    assert sections[0]["text"] == "疑似实际控制人：梁建章"
    assert sections[2]["rows"][0][1] == "58.8%"


def test_v32_overview_recognizer_splits_email_relation() -> None:
    sections = sections_from_snapshot(
        MODULE_BY_KEY["overview"],
        {
            "keyValues": [
                ["电子邮箱", "service@example.com；同邮箱企业 3 家"],
                ["企业名称", "上海携程金融信息服务有限公司"],
            ],
            "tables": [],
        },
    )

    rows = dict(sections[0]["rows"])
    assert rows["电子邮箱"] == "service@example.com"
    assert rows["同邮箱企业"] == "3家"


def test_v32_entity_and_avatar_cleanup_is_applied_by_semantic_column() -> None:
    executive = sections_from_snapshot(
        MODULE_BY_KEY["executives"],
        {"tables": [{"title": "高管信息", "rows": [["姓名", "职务"], ["兰 兰云", "董事"]]}]},
    )
    investment = sections_from_snapshot(
        MODULE_BY_KEY["investments"],
        {
            "tables": [
                {
                    "title": "对外投资企业",
                    "rows": [
                        ["企业名称", "投资比例", "企业状态", "行业"],
                        ["上 上海测试有限公司民企民企", "100%", "存续", "软件"],
                    ],
                }
            ]
        },
    )

    assert executive[0]["rows"][0][0] == "兰云"
    assert investment[0]["rows"][0][0] == "上海测试有限公司（民企）"


def test_v32_business_changes_use_separate_before_and_after_columns() -> None:
    sections = sections_from_snapshot(
        MODULE_BY_KEY["business_changes"],
        {
            "tables": [
                {
                    "title": "工商变更",
                    "rows": [
                        ["序号", "变更时间", "变更项目", "变更内容"],
                        ["1", "2026-09-01", "住所", "旧地址", "新地址"],
                    ],
                }
            ]
        },
    )

    assert sections[0]["columns"][:5] == ["序号", "变更时间", "变更项目", "变更前", "变更后"]
    assert sections[0]["rows"][0][3:] == ["旧地址", "新地址"]
    assert _business_change_needs_review(sections) is False


def test_v32_business_changes_reject_systemically_empty_after_column() -> None:
    sections = [
        {
            "title": "工商变更",
            "columns": ["序号", "变更时间", "变更项目", "变更前", "变更后"],
            "rows": [["1", "2026-09-01", "住所", "旧地址", ""]],
        }
    ]

    assert _business_change_needs_review(sections) is True


def test_target_company_is_queued_without_waiting_for_legacy_ingestion(monkeypatch) -> None:
    from app.api.routes import companies as companies_route

    monkeypatch.setattr(companies_route.get_settings(), "company_profile_worker_enabled", True)

    async def unexpected_refresh(*args, **kwargs):
        raise AssertionError("目标企业不应等待原有公开源刷新")

    monkeypatch.setattr(
        companies_route,
        "refresh_company_akshare_profile",
        unexpected_refresh,
    )
    monkeypatch.setattr(
        companies_route,
        "refresh_company_qichacha_profile",
        unexpected_refresh,
    )
    db = database()

    response = asyncio.run(
        companies_route.search_and_ingest_company(
            CompanySearchIngestRequest(name="上海携程金融信息服务有限公司"),
            db,
        )
    )

    assert response.company_profile is not None
    assert response.company_profile["run"]["status"] == "queued"
    assert response.ingestion_run is None


def test_safe_company_alias_is_canonicalized_before_creation(monkeypatch) -> None:
    from app.api.routes import companies as companies_route

    monkeypatch.setattr(companies_route.get_settings(), "company_profile_worker_enabled", True)

    async def unexpected_refresh(*args, **kwargs):
        raise AssertionError("安全别名应直接进入企业全景队列")

    monkeypatch.setattr(companies_route, "refresh_company_akshare_profile", unexpected_refresh)
    monkeypatch.setattr(companies_route, "refresh_company_qichacha_profile", unexpected_refresh)
    db = database()

    response = asyncio.run(
        companies_route.search_and_ingest_company(
            CompanySearchIngestRequest(name=" 携程金融信息服务有限公司 "),
            db,
        )
    )

    assert response.company.name == "上海携程金融信息服务有限公司"
    assert response.resolved_from == "携程金融信息服务有限公司"
    assert response.company_profile is not None
    assert db.query(Company).count() == 1


def test_short_company_alias_requires_confirmation_without_creating_company() -> None:
    from app.api.routes import companies as companies_route

    db = database()
    with pytest.raises(HTTPException) as exc_info:
        asyncio.run(
            companies_route.search_and_ingest_company(
                CompanySearchIngestRequest(name="携程金融"),
                db,
            )
        )

    assert exc_info.value.status_code == 409
    assert exc_info.value.detail["code"] == "company_confirmation_required"
    assert exc_info.value.detail["suggestion"]["name"] == "上海携程金融信息服务有限公司"
    assert db.query(Company).count() == 0


def test_company_alias_registry_is_deliberately_bounded() -> None:
    assert resolve_profile_company_name("上海 携程金融信息服务有限公司") == (
        "上海携程金融信息服务有限公司"
    )
    assert resolve_profile_company_name("携程金融") is None
    assert resolve_profile_company_name("携程") is None
    assert profile_company_suggestion("携程金融")["requires_confirmation"] is True
    assert profile_company_suggestion("携程金融信息")["requires_confirmation"] is True
    assert profile_company_suggestion("携程") is None


def test_existing_safe_alias_record_is_migrated_to_canonical_name(monkeypatch) -> None:
    from app.api.routes import companies as companies_route

    monkeypatch.setattr(companies_route.get_settings(), "company_profile_worker_enabled", True)

    async def unexpected_refresh(*args, **kwargs):
        raise AssertionError("旧别名记录应直接迁移并进入队列")

    monkeypatch.setattr(companies_route, "refresh_company_akshare_profile", unexpected_refresh)
    monkeypatch.setattr(companies_route, "refresh_company_qichacha_profile", unexpected_refresh)
    db = database()
    legacy = Company(name="携程金融信息服务有限公司", company_profile={})
    db.add(legacy)
    db.commit()

    response = asyncio.run(
        companies_route.search_and_ingest_company(
            CompanySearchIngestRequest(name="上海携程金融信息服务有限公司"),
            db,
        )
    )

    assert response.company.id == legacy.id
    assert response.company.name == "上海携程金融信息服务有限公司"
    assert db.query(Company).count() == 1


def test_profile_target_does_not_fuzzy_match_a_different_company(monkeypatch) -> None:
    from app.api.routes import companies as companies_route

    monkeypatch.setattr(companies_route.get_settings(), "company_profile_worker_enabled", True)

    async def unexpected_refresh(*args, **kwargs):
        raise AssertionError("目标企业应直接进入企业全景队列")

    monkeypatch.setattr(companies_route, "refresh_company_akshare_profile", unexpected_refresh)
    monkeypatch.setattr(companies_route, "refresh_company_qichacha_profile", unexpected_refresh)
    db = database()
    db.add(Company(name="上海携程金融信息服务有限公司分公司", company_profile={}))
    db.commit()

    response = asyncio.run(
        companies_route.search_and_ingest_company(
            CompanySearchIngestRequest(name="上海携程金融信息服务有限公司"),
            db,
        )
    )

    assert response.company.name == "上海携程金融信息服务有限公司"
    assert db.query(Company).count() == 2


def test_company_profile_contract_has_three_reusable_page_families() -> None:
    assert {spec.page for spec in MODULE_SPECS} == {"overview", "creditData", "monitor"}
    assert all(spec.container_selectors for spec in MODULE_SPECS)
    assert all(spec.heading_aliases for spec in MODULE_SPECS)


def test_worker_keeps_running_when_failure_callback_is_temporarily_unreachable(capsys) -> None:
    class UnavailableApi:
        def post(self, path, payload):
            raise WorkerApiError("临时无法连接")

    agent = CompanyProfileAgent(
        UnavailableApi(),
        BrowserConfig(user_data_dir=Path(".")),
        worker_id="windows-test",
    )

    agent._fail("run-id", "failed", "network", "test", {})

    assert "company_profile_failure_report_deferred" in capsys.readouterr().out


def test_collector_reuses_one_tab_per_page_family(monkeypatch) -> None:
    from enterprise_sentinel.company_profile.collector import CompanyProfileCollector

    class FakePage:
        def __init__(self, number: int) -> None:
            self.number = number
            self.closed = False

        def close(self) -> None:
            self.closed = True

    class FakeBrowser:
        def __init__(self) -> None:
            self.pages: list[FakePage] = []

        def new_tab(self) -> FakePage:
            page = FakePage(len(self.pages) + 1)
            self.pages.append(page)
            return page

    class FakeEngine:
        def __init__(self) -> None:
            self.browser = FakeBrowser()

    collector = object.__new__(CompanyProfileCollector)
    collector.engine = FakeEngine()
    collector.retry_times = 2
    observed_pages: dict[str, set[int]] = {}

    def fake_collect_module(
        page,
        spec,
        company_name,
        company_code,
        *,
        force_reload,
        fallback_scope,
        retry_count,
    ):
        observed_pages.setdefault(spec.page, set()).add(page.number)
        timing = {
            "navigation_ms": 1,
            "target_wait_ms": 1,
            "extract_ms": 1,
            "pagination_wait_ms": 0,
            "total_ms": 3,
            "retry_count": retry_count,
            "fallback_level": "L0_TARGET",
        }
        module = {
            "key": spec.key,
            "title": spec.title,
            "status": "PASS",
            "source_url": spec.url(company_code),
            "captured_at": datetime.now(timezone.utc).isoformat(),
            "page_count": 1,
            "row_count": 1,
            "sections": [],
            "timing": timing,
        }
        return module, {"source_url": module["source_url"], "pages": [], "timing": timing}

    monkeypatch.setattr(collector, "_collect_module", fake_collect_module)

    result = collector.collect("上海携程金融信息服务有限公司", company_code="cached-code")

    assert len(collector.engine.browser.pages) == 3
    assert all(len(page_ids) == 1 for page_ids in observed_pages.values())
    assert all(page.closed for page in collector.engine.browser.pages)
    assert result["raw_data"]["performance"]["page_family_count"] == 3
    assert set(result["raw_data"]["performance"]["module_timings"]) == {
        spec.key for spec in MODULE_SPECS
    }


def test_collector_reports_region_access_denial() -> None:
    from enterprise_sentinel.company_profile.collector import CompanyProfileCollector

    class FakeEngine:
        @staticmethod
        def _raise_if_login_redirect(page) -> None:
            return None

    class FakePage:
        text = ""

        @staticmethod
        def run_js(script):
            return {
                "captcha": False,
                "permissionDenied": False,
                "regionDenied": True,
            }

    collector = object.__new__(CompanyProfileCollector)
    collector.engine = FakeEngine()

    with pytest.raises(PermissionError, match="中国大陆网络"):
        collector._check_access(FakePage())


def test_windows_powershell_scripts_use_utf8_bom() -> None:
    scripts = sorted((ROOT / "deploy" / "windows").glob("*.ps1"))

    assert scripts
    assert all(script.read_bytes().startswith(b"\xef\xbb\xbf") for script in scripts)


def test_windows_worker_uses_dedicated_live_edge_profile() -> None:
    runner = (ROOT / "deploy" / "windows" / "run-worker.ps1").read_text(
        encoding="utf-8-sig"
    )

    assert "'--use-live-profile'" in runner
    assert "'--clone-root-dir'" not in runner


def test_browser_login_message_names_edge() -> None:
    engine = object.__new__(CrawlerEngine)
    engine.browser_config = BrowserConfig(
        user_data_dir=Path("profile"), browser_path=r"C:\Program Files\Microsoft\Edge\msedge.exe"
    )

    assert engine._browser_label() == "Microsoft Edge"
