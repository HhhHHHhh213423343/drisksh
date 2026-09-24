from __future__ import annotations

from pathlib import Path
import sys
from datetime import datetime, timezone

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))

from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import Session

from app.config import get_settings
from app.db.base import Base
from app.models import (
    AnalysisRun,
    Company,
    ExtractedFact,
    MacroIndicatorPoint,
    MonthlyReport,
    MonitoringRule,
    RiskEvent,
    RuleEvaluation,
    SourceDocument,
)
from app.services.document_ingestion import METRIC_ALIASES, parse_path, safe_filename, validate_extension
from app.services.monthly_reports import (
    SECTION_ORDER,
    export_monthly_report_xlsx,
    generate_monthly_report,
)
from app.services.rule_engine import evaluate_rules
from app.services.rule_seed import RULE_CONDITIONS, ensure_company_rule_assignments, ensure_default_rule_set


def database() -> Session:
    engine = create_engine("sqlite:///:memory:", future=True)
    Base.metadata.create_all(engine)
    return Session(engine)


def test_rule_seed_contains_all_79_rules() -> None:
    db = database()
    rule_set = ensure_default_rule_set(db)
    assert rule_set.version == "2026.1"
    assert db.scalar(select(func.count(MonitoringRule.id))) == 79
    assert db.scalar(select(MonitoringRule).where(MonitoringRule.code == 32)).condition
    required_metrics = {
        item["metric"]
        for condition in RULE_CONDITIONS.values()
        for group in ("all", "any")
        for item in condition.get(group, [])
    }
    assert required_metrics <= set(METRIC_ALIASES)


def test_spreadsheet_parser_extracts_period_facts(tmp_path: Path) -> None:
    path = tmp_path / "finance.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["指标", "2025-06", "2026-06"])
    sheet.append(["营业收入", 100, 80])
    sheet.append(["净利润", 10, -2])
    workbook.save(path)

    result = parse_path(path, ".xlsx")
    facts = {(fact.metric_code, fact.period, fact.value_numeric) for fact in result.facts}
    assert ("revenue", "2025-06", 100.0) in facts
    assert ("revenue", "2026-06", 80.0) in facts
    assert ("net_profit", "2026-06", -2.0) in facts


def test_upload_filename_and_macro_file_are_rejected() -> None:
    assert safe_filename("../../财务报告.xlsx") == "财务报告.xlsx"
    try:
        validate_extension("含宏文件.xlsm")
    except ValueError as exc:
        assert "仅支持" in str(exc)
    else:
        raise AssertionError("宏文件不应进入解析流程")


def test_rule_engine_distinguishes_hit_and_missing_data(monkeypatch) -> None:
    monkeypatch.setenv("LLM_BASE_URL", "")
    monkeypatch.setenv("LLM_MODEL", "")
    get_settings.cache_clear()
    db = database()
    company = Company(name="测试商业保理有限公司")
    db.add(company)
    db.flush()
    ensure_company_rule_assignments(db, company.id)
    document = SourceDocument(
        company_id=company.id,
        original_filename="finance.xlsx",
        extension=".xlsx",
        storage_path="/tmp/finance.xlsx",
        sha256="a" * 64,
        size_bytes=10,
        status="parsed",
    )
    db.add(document)
    db.flush()
    db.add_all(
        [
            ExtractedFact(
                document_id=document.id,
                company_id=company.id,
                metric_code="revenue",
                label="营业收入",
                value_numeric=100,
                value_text="100",
                period="2025-06",
                confidence=1,
                status="confirmed",
            ),
            ExtractedFact(
                document_id=document.id,
                company_id=company.id,
                metric_code="revenue",
                label="营业收入",
                value_numeric=80,
                value_text="80",
                period="2026-06",
                confidence=1,
                status="confirmed",
            ),
        ]
    )
    run = AnalysisRun(
        company_id=company.id,
        status="running",
        requested_period="2026-06",
        current_stage="evaluate_rules",
    )
    db.add(run)
    db.commit()

    result = evaluate_rules(db, run.id)
    assert result["rule_count"] == 79
    rule_32 = db.execute(
        select(RuleEvaluation)
        .join(MonitoringRule, MonitoringRule.id == RuleEvaluation.rule_id)
        .where(RuleEvaluation.analysis_run_id == run.id)
        .where(MonitoringRule.code == 32)
    ).scalar_one()
    rule_33 = db.execute(
        select(RuleEvaluation)
        .join(MonitoringRule, MonitoringRule.id == RuleEvaluation.rule_id)
        .where(RuleEvaluation.analysis_run_id == run.id)
        .where(MonitoringRule.code == 33)
    ).scalar_one()
    assert rule_32.status == "hit"
    assert rule_33.status == "insufficient_data"


def test_rule_engine_limits_events_to_requested_month(monkeypatch) -> None:
    monkeypatch.setenv("LLM_BASE_URL", "")
    monkeypatch.setenv("LLM_MODEL", "")
    get_settings.cache_clear()
    db = database()
    company = Company(name="期间测试公司")
    db.add(company)
    db.flush()
    ensure_company_rule_assignments(db, company.id)
    db.add_all(
        [
            RiskEvent(
                company_id=company.id,
                category="legal",
                severity="important",
                title="7月新增诉讼",
                content="公司发生重大诉讼，已发布开庭公告。",
                source_url="https://example.com/july",
                source_name="测试来源",
                occurred_at=datetime(2026, 7, 12, tzinfo=timezone.utc),
            ),
            RiskEvent(
                company_id=company.id,
                category="legal",
                severity="important",
                title="8月新增诉讼",
                content="公司发生重大诉讼，已发布开庭公告。",
                source_url="https://example.com/august",
                source_name="测试来源",
                occurred_at=datetime(2026, 8, 2, tzinfo=timezone.utc),
            ),
        ]
    )
    run = AnalysisRun(
        company_id=company.id,
        status="running",
        requested_period="2026-07",
        current_stage="evaluate_rules",
    )
    db.add(run)
    db.commit()
    evaluate_rules(db, run.id)
    evaluation = db.execute(
        select(RuleEvaluation)
        .join(MonitoringRule, MonitoringRule.id == RuleEvaluation.rule_id)
        .where(RuleEvaluation.analysis_run_id == run.id)
        .where(MonitoringRule.code == 73)
    ).scalar_one()
    assert evaluation.status == "hit"
    assert [item["title"] for item in evaluation.evidence] == ["7月新增诉讼"]


def test_monthly_report_generates_versions_and_xlsx(monkeypatch) -> None:
    monkeypatch.setenv("LLM_BASE_URL", "")
    monkeypatch.setenv("LLM_MODEL", "")
    get_settings.cache_clear()
    monkeypatch.setattr(get_settings(), "demo_report_as_of_date", "2026-07-09", raising=False)
    db = database()
    company = Company(name="月报测试公司")
    db.add(company)
    db.flush()
    ensure_company_rule_assignments(db, company.id)
    document = SourceDocument(
        company_id=company.id,
        original_filename="monthly.xlsx",
        extension=".xlsx",
        storage_path="/tmp/monthly.xlsx",
        sha256="b" * 64,
        size_bytes=10,
        status="parsed",
    )
    db.add(document)
    db.flush()
    db.add_all(
        [
            ExtractedFact(
                document_id=document.id,
                company_id=company.id,
                metric_code="revenue",
                label="营业收入",
                value_numeric=100,
                value_text="100",
                period="2026-06",
                confidence=1,
                status="confirmed",
            ),
            ExtractedFact(
                document_id=document.id,
                company_id=company.id,
                metric_code="revenue",
                label="营业收入",
                value_numeric=120,
                value_text="120",
                period="2026-07",
                confidence=1,
                status="confirmed",
            ),
        ]
    )
    db.add(
        MacroIndicatorPoint(
            indicator_code="manufacturing_pmi",
            indicator_name="制造业 PMI",
            period="2026-07",
            value=49.2,
            unit="点",
            frequency="月度",
            region_scope="全国",
            source_name="国家统计局",
            source_url="https://data.stats.gov.cn/",
        )
    )
    run = AnalysisRun(
        company_id=company.id,
        status="running",
        requested_period="2026-07",
        current_stage="evaluate_rules",
    )
    db.add(run)
    db.commit()
    evaluate_rules(db, run.id)
    run.status = "completed"
    db.add(run)
    report = MonthlyReport(
        company_id=company.id,
        analysis_run_id=run.id,
        period="2026-07",
        version=1,
        status="generating",
        title="月报测试公司2026-07投后监测报告",
        created_by="tester",
    )
    db.add(report)
    db.commit()

    generated = generate_monthly_report(db, report.id)
    assert generated.status == "draft"
    assert list(generated.sections) == SECTION_ORDER
    assert "较上月增长20.00%" in generated.sections["finance"]
    assert "2026-07制造业 PMI49.2点" in generated.sections["external"]
    assert generated.evidence_snapshot["as_of_date"] == "2026-07-09"
    assert generated.evidence_snapshot["period_status"] == "month_to_date"
    content = export_monthly_report_xlsx(db, generated)
    workbook = load_workbook(filename=__import__("io").BytesIO(content), read_only=True)
    assert workbook.sheetnames == ["风险项目", "规则评估", "证据索引"]
    assert workbook["规则评估"].max_row == 80
    evidence_rows = list(workbook["证据索引"].iter_rows(values_only=True))
    assert any("https://data.stats.gov.cn/" in str(row[5]) for row in evidence_rows)
