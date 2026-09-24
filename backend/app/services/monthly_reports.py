from __future__ import annotations

import io
import calendar
import re
from collections import defaultdict
from datetime import date, datetime, timezone
from typing import Any

from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.models import (
    AnalysisRun,
    Company,
    ExtractedFact,
    MacroIndicatorPoint,
    MacroIndustryEvent,
    MonthlyReport,
    MonitoringRule,
    ReportVersion,
    RiskEvent,
    RuleEvaluation,
)
from app.config import get_settings
from app.services.llm import LLMService


SECTION_ORDER = [
    "risk_overview",
    "actions",
    "external",
    "operations",
    "finance",
    "governance",
    "internal_control",
    "legal",
    "brand",
    "sentiment",
    "other",
]

SECTION_LABELS = {
    "risk_overview": "风险概况",
    "actions": "行动建议",
    "external": "外部风险",
    "operations": "业务及经营风险",
    "finance": "财务风险",
    "governance": "公司治理风险",
    "internal_control": "内控风险",
    "legal": "法律风险",
    "brand": "品牌风险",
    "sentiment": "舆情风险",
    "other": "其他风险",
}

STYLE_GUIDE = (
    "文风模仿风险管理台账月报：结论先行，先用一小段概括本月整体特征与较上期变化，再分板块叙述；"
    "每个板块用连贯段落陈述事实与判断，需要评估影响时以“风险分析：”引出；"
    "没有发现风险时直接写“本期暂未发现相关风险”；"
    "对有数据的财务经营事项写清本期值、上期值及变化原因；对风险事项写清主体、事件、影响和待跟踪动作；"
    "面向客户的正文中严禁出现“规则”“命中N项”“监测要点”“insufficient_data”等内部评估术语；"
    "数据不足时写明尚缺何种材料、后续如何补充，不作风险推断；"
    "行动建议使用“跟踪、评估、核验、清收、完善、审慎”等动词，写成可执行的具体动作；"
    "不要用宣传性、夸张性或模板化语言。"
)

CATEGORY_SECTIONS = {
    "外部风险": "external",
    "业务及经营风险": "operations",
    "财务风险": "finance",
    "公司治理风险": "governance",
    "内控风险": "internal_control",
    "法律风险": "legal",
    "其他风险": "other",
}

FINANCE_METRIC_LABELS = {
    "revenue": "营业收入",
    "net_profit": "净利润",
    "net_assets": "净资产",
    "cash": "货币资金",
    "receivables": "应收款项",
    "operating_cashflow": "经营活动现金流量净额",
}

MACRO_INDICATOR_ORDER = (
    "gdp_yoy",
    "cpi_yoy",
    "ppi_yoy",
    "manufacturing_pmi",
    "non_manufacturing_pmi",
    "lpr_1y",
    "m2_yoy",
)


def _section_for_rule(rule: MonitoringRule) -> str:
    if rule.category == "品牌舆情风险":
        return "sentiment" if rule.code == 76 else "brand"
    return CATEGORY_SECTIONS.get(rule.category, "other")


def _evaluations(db: Session, run_id: Any) -> list[tuple[RuleEvaluation, MonitoringRule]]:
    return list(
        db.execute(
            select(RuleEvaluation, MonitoringRule)
            .join(MonitoringRule, MonitoringRule.id == RuleEvaluation.rule_id)
            .where(RuleEvaluation.analysis_run_id == run_id)
            .order_by(MonitoringRule.code)
        ).all()
    )


def _previous_month(period: str) -> str:
    year, month = (int(part) for part in period.split("-"))
    return f"{year - 1}-12" if month == 1 else f"{year}-{month - 1:02d}"


def _report_as_of_date(period: str) -> date:
    configured = get_settings().demo_report_as_of_date
    if configured:
        try:
            parsed = date.fromisoformat(configured)
            if parsed.strftime("%Y-%m") == period:
                return parsed
        except ValueError:
            pass
    today = datetime.now().date()
    if today.strftime("%Y-%m") == period:
        return today
    year, month = (int(part) for part in period.split("-"))
    return date(year, month, calendar.monthrange(year, month)[1])


def _period_rank(value: str) -> tuple[int, int]:
    text = str(value or "")
    month_match = re.search(r"(20\d{2})[-/.](0?[1-9]|1[0-2])", text)
    if month_match:
        return int(month_match.group(1)), int(month_match.group(2))
    quarter_match = re.search(r"(20\d{2}).*?[Qq季](?:度)?[- ]?([1-4])", text)
    if not quarter_match:
        quarter_match = re.search(r"(20\d{2})[- ]?[Qq]([1-4])", text)
    if quarter_match:
        return int(quarter_match.group(1)), int(quarter_match.group(2)) * 3
    year_match = re.search(r"(20\d{2})", text)
    return (int(year_match.group(1)), 12) if year_match else (0, 0)


def _macro_context(
    db: Session, period: str, company_id: Any = None
) -> list[dict[str, Any]]:
    cutoff = _period_rank(period)
    rows = list(
        db.execute(
            select(MacroIndicatorPoint)
            .where(MacroIndicatorPoint.region_scope == "全国")
            .order_by(MacroIndicatorPoint.collected_at.desc())
            .limit(1000)
        ).scalars()
    )
    latest: dict[str, MacroIndicatorPoint] = {}
    for row in rows:
        rank = _period_rank(row.period)
        if rank > cutoff or row.indicator_code in latest:
            continue
        latest[row.indicator_code] = row
    result: list[dict[str, Any]] = []
    for code in MACRO_INDICATOR_ORDER:
        row = latest.get(code)
        if not row:
            continue
        result.append(
            {
                "kind": "indicator",
                "evidence_id": f"macro-indicator-{row.id}",
                "code": row.indicator_code,
                "label": row.indicator_name,
                "period": row.period,
                "value": row.value,
                "unit": row.unit,
                "source_name": row.source_name,
                "source_url": row.source_url,
            }
        )

    as_of = _report_as_of_date(period)
    end = datetime(as_of.year, as_of.month, as_of.day, 23, 59, 59, tzinfo=timezone.utc)
    events = list(
        db.execute(
            select(MacroIndustryEvent)
            .where(MacroIndustryEvent.published_at <= end)
            .order_by(MacroIndustryEvent.published_at.desc())
            .limit(12)
        ).scalars()
    )
    result.extend(
        {
            "kind": "event",
            "evidence_id": f"macro-event-{event.id}",
            "dimension": event.dimension,
            "label": event.title,
            "period": event.published_at.date().isoformat(),
            "summary": event.summary,
            "source_name": event.source_name,
            "source_url": event.source_url,
        }
        for event in events
    )
    if company_id:
        company_events = list(
            db.execute(
                select(RiskEvent)
                .where(RiskEvent.company_id == company_id)
                .where(RiskEvent.occurred_at <= end)
                .order_by(RiskEvent.occurred_at.desc())
                .limit(12)
            ).scalars()
        )
        result.extend(
            {
                "kind": "company_event",
                "evidence_id": f"risk-event-{event.id}",
                "category": event.category,
                "severity": event.severity,
                "sentiment": event.sentiment,
                "label": event.title,
                "period": event.occurred_at.date().isoformat() if event.occurred_at else "",
                "summary": event.content,
                "source_name": event.source_name,
                "source_url": event.source_url,
            }
            for event in company_events
        )
    return result


def _macro_sentence(context: list[dict[str, Any]], as_of: date) -> str:
    indicators = [item for item in context if item.get("kind") == "indicator"]
    if not indicators:
        return (
            f"截至{as_of.isoformat()}，尚未取得可追溯的当期国家宏观指标，"
            "外部环境结论仅依据已入库事件，不对缺失数据作推断。"
        )
    parts = [
        f"{item['period']}{item['label']}{float(item['value']):g}{item.get('unit') or ''}"
        for item in indicators
    ]
    sources = "、".join(
        dict.fromkeys(str(item.get("source_name") or "") for item in indicators)
    )
    return f"截至{as_of.isoformat()}，{sources}公开指标显示：" + "；".join(parts) + "。"


def _monthly_metric_context(
    db: Session, report: MonthlyReport
) -> list[dict[str, Any]]:
    previous = _previous_month(report.period)
    facts = list(
        db.execute(
            select(ExtractedFact)
            .where(ExtractedFact.company_id == report.company_id)
            .where(ExtractedFact.status == "confirmed")
            .where(ExtractedFact.metric_code.in_(FINANCE_METRIC_LABELS))
            .where(ExtractedFact.period.in_([report.period, previous]))
            .order_by(ExtractedFact.created_at.asc())
        ).scalars()
    )
    latest: dict[tuple[str, str], ExtractedFact] = {
        (fact.metric_code, fact.period): fact for fact in facts
    }
    result = []
    for code, label in FINANCE_METRIC_LABELS.items():
        current = latest.get((code, report.period))
        prior = latest.get((code, previous))
        if not current or current.value_numeric is None:
            continue
        change = None
        if prior and prior.value_numeric not in (None, 0):
            change = float(current.value_numeric) / float(prior.value_numeric) - 1
        result.append(
            {
                "metric_code": code,
                "label": label,
                "period": report.period,
                "value": current.value_numeric,
                "value_text": current.value_text,
                "unit": current.unit,
                "previous_period": previous,
                "previous_value": prior.value_numeric if prior else None,
                "month_over_month": change,
                "evidence_id": f"fact-{current.id}",
                "source_locator": current.source_locator,
                "document_id": str(current.document_id),
            }
        )
    existing_codes = {item["metric_code"] for item in result}
    company = db.get(Company, report.company_id)
    profile_rows = (
        ((company.company_profile or {}).get("akshare_profile") or {}).get(
            "financial_abstract"
        )
        if company
        else []
    ) or []
    aliases = {
        "revenue": ("营业总收入", "营业收入"),
        "net_profit": ("归母净利润", "净利润"),
        "net_assets": ("归母股东权益", "股东权益", "净资产"),
        "operating_cashflow": ("经营现金净流量", "经营活动产生的现金流量净额"),
    }
    report_period_key = report.period.replace("-", "") + "99"
    for code, names in aliases.items():
        if code in existing_codes:
            continue
        row = next(
            (
                item
                for item in profile_rows
                if any(
                    name in str(item.get("指标") or item.get("项目") or "")
                    for name in names
                )
            ),
            None,
        )
        if not row:
            continue
        values = []
        for key, value in row.items():
            key_text = str(key)
            if not re.fullmatch(r"20\d{6}", key_text) or key_text > report_period_key:
                continue
            try:
                number = float(str(value).replace(",", ""))
            except (TypeError, ValueError):
                continue
            values.append((key_text, number))
        if not values:
            continue
        latest_period, value = sorted(values)[-1]
        result.append(
            {
                "metric_code": code,
                "label": f"{FINANCE_METRIC_LABELS[code]}（最新披露期{latest_period}）",
                "period": latest_period,
                "value": value,
                "value_text": str(value),
                "unit": "元",
                "previous_period": "",
                "previous_value": None,
                "month_over_month": None,
                "evidence_id": f"akshare-{code}-{latest_period}",
                "source_locator": f"company_profile.akshare_profile.financial_abstract.{latest_period}",
            }
        )
    return result


def _metric_sentence(metrics: list[dict[str, Any]]) -> str:
    parts = []
    for metric in metrics:
        value = float(metric["value"])
        unit = metric.get("unit") or ""
        if unit == "元":
            display = f"{value / 100_000_000:,.2f}亿元"
        elif unit == "%":
            display = f"{value * 100:.2f}%"
        else:
            display = metric.get("value_text") or f"{value:,.2f}"
        change = metric.get("month_over_month")
        comparison = (
            f"，较上月{'增长' if change >= 0 else '下降'}{abs(change) * 100:.2f}%"
            if change is not None
            else "，上月同口径数据尚未确认"
        )
        parts.append(f"{metric['label']}为{display}{comparison}")
    return "；".join(parts) + "。" if parts else ""


def _deterministic_sections(
    company: Company,
    period: str,
    rows: list[tuple[RuleEvaluation, MonitoringRule]],
    monthly_metrics: list[dict[str, Any]],
    macro_context: list[dict[str, Any]],
    as_of_date: date,
) -> tuple[str, dict[str, str], dict[str, list[str]]]:
    grouped: dict[str, list[tuple[RuleEvaluation, MonitoringRule]]] = defaultdict(list)
    evidence_by_section: dict[str, list[str]] = defaultdict(list)
    for evaluation, rule in rows:
        section = _section_for_rule(rule)
        grouped[section].append((evaluation, rule))
        evidence_by_section[section].extend(
            str(item.get("id"))
            for item in (evaluation.evidence or [])
            if item.get("id")
        )

    hits = [(evaluation, rule) for evaluation, rule in rows if evaluation.status == "hit"]
    gaps = [
        (evaluation, rule)
        for evaluation, rule in rows
        if evaluation.status == "insufficient_data"
    ]
    important = [item for item in hits if item[0].severity in {"important", "high", "critical"}]
    summary = (
        f"{company.name}{period}投后监测（截至{as_of_date.isoformat()}）已完成{len(rows)}项监测要点核验，"
        f"发现{len(hits)}项需关注事项（其中重要事项{len(important)}项）；"
        f"另有{len(gaps)}项因资料不足，待补充材料后进一步核验。"
    )
    sections: dict[str, str] = {}
    sections["risk_overview"] = summary
    if hits:
        actions = []
        for index, (evaluation, rule) in enumerate(hits[:6], start=1):
            actions.append(
                f"{index}. 针对“{rule.monitoring_point}”核验最新进展，"
                f"明确责任人与完成时点，落实风险缓释措施。"
            )
        sections["actions"] = " ".join(actions)
    elif gaps:
        sections["actions"] = "1. 补充本期财务、经营及治理资料，完成相关事项的复核。"
    else:
        sections["actions"] = "持续按月更新公开信息和公司经营材料。"

    for section in SECTION_ORDER[2:]:
        items = grouped.get(section, [])
        section_hits = [(evaluation, rule) for evaluation, rule in items if evaluation.status == "hit"]
        section_gaps = [
            (evaluation, rule)
            for evaluation, rule in items
            if evaluation.status == "insufficient_data"
        ]
        if section_hits:
            paragraphs = [
                f"本期该板块发现{len(section_hits)}项需关注事项。"
            ]
            for evaluation, rule in section_hits:
                paragraphs.append(
                    f"{rule.monitoring_point}：{evaluation.rationale}"
                )
            if section_gaps:
                paragraphs.append(
                    f"另有{len(section_gaps)}项监测事项因资料不足暂缓结论，需继续补充原始材料。"
                )
            sections[section] = "\n\n".join(paragraphs)
        elif section_gaps:
            sections[section] = (
                f"本期该板块所需监测材料尚未齐备，共有{len(section_gaps)}项监测事项暂缓结论，"
                "不能据此认定不存在风险，待补充原始材料后进一步核验。"
            )
        else:
            sections[section] = "本期暂未发现相关风险。"
    metric_summary = _metric_sentence(monthly_metrics)
    if metric_summary:
        sections["finance"] = metric_summary + "\n\n" + sections["finance"]
    macro_summary = _macro_sentence(macro_context, as_of_date)
    sections["external"] = macro_summary + "\n\n" + sections["external"]
    evidence_by_section["external"].extend(
        str(item["evidence_id"])
        for item in macro_context
        if item.get("evidence_id")
    )
    return summary, sections, {
        key: list(dict.fromkeys(values)) for key, values in evidence_by_section.items()
    }


def _previous_report(db: Session, report: MonthlyReport) -> MonthlyReport | None:
    return db.execute(
        select(MonthlyReport)
        .where(MonthlyReport.company_id == report.company_id)
        .where(MonthlyReport.status == "approved")
        .where(MonthlyReport.period < report.period)
        .order_by(MonthlyReport.period.desc(), MonthlyReport.version.desc())
        .limit(1)
    ).scalar_one_or_none()


def _llm_sections(
    db: Session,
    *,
    report: MonthlyReport,
    company: Company,
    rows: list[tuple[RuleEvaluation, MonitoringRule]],
    base_summary: str,
    base_sections: dict[str, str],
    evidence_by_section: dict[str, list[str]],
    monthly_metrics: list[dict[str, Any]],
    macro_context: list[dict[str, Any]],
    as_of_date: date,
) -> tuple[str, dict[str, str], dict[str, Any], str]:
    service = LLMService()
    if not service.is_configured:
        return base_summary, base_sections, {}, "deterministic"
    previous = _previous_report(db, report)
    hits = [
        {
            "rule_code": rule.code,
            "section": _section_for_rule(rule),
            "monitoring_point": rule.monitoring_point,
            "assessment_standard": rule.assessment_standard,
            "severity": evaluation.severity,
            "rationale": evaluation.rationale,
            "evidence_ids": [
                item.get("id") for item in (evaluation.evidence or []) if item.get("id")
            ],
            "facts": evaluation.facts or [],
        }
        for evaluation, rule in rows
        if evaluation.status == "hit"
    ]
    gap_counts: dict[str, int] = defaultdict(int)
    for evaluation, rule in rows:
        if evaluation.status == "insufficient_data":
            gap_counts[_section_for_rule(rule)] += 1
    previous_style = {
        key: str((previous.sections or {}).get(key) or "")[:1200]
        for key in SECTION_ORDER
        if previous and (previous.sections or {}).get(key)
    }
    valid_evidence = {
        value for values in evidence_by_section.values() for value in values
    }
    valid_evidence.update(
        str(metric["evidence_id"])
        for metric in monthly_metrics
        if metric.get("evidence_id")
    )
    try:
        result = service.complete_json(
            db,
            purpose=f"monthly-report:{company.id}:{report.period}",
            system_prompt=(
                "你是投后风险管理月报撰写助手，最终读者是投资机构管理层，输出将直接进入正式风险管理台账。"
                "文风模仿风险管理台账月报：结论先行，先用一小段概括本月整体特征与较上期变化；"
                "各板块用连贯段落陈述事实与判断，需要评估影响时以“风险分析：”引出；"
                "没有发现风险时直接写“本期暂未发现相关风险”。"
                "严禁在正文中出现“规则”“命中N项”“监测要点”“规则编号”“insufficient_data”等内部评估术语，"
                "也不要罗列评估数量，应把评估结论转化为对公司经营、财务或合规状况的具体影响描述。"
                "external（外部风险）板块按“行业监管态势”“同业动态”“风险分析：”三段组织，"
                "行业监管态势写监管政策与会议动向及其影响，同业动态写行业结构变化与竞争格局。"
                + STYLE_GUIDE
                + "只能使用输入事实与评估结论，不得创造数字、案件或来源；数据不足必须明确说明，不得推断。"
                "返回JSON对象：{summary:string,sections:{risk_overview,actions,external,operations,finance,"
                "governance,internal_control,legal,brand,sentiment,other},evidence_ids_by_section:{...}}。"
                "正文中不要声称本报告已获审批。"
            ),
            user_payload={
                "company": company.name,
                "period": report.period,
                "report_as_of_date": as_of_date.isoformat(),
                "base_summary": base_summary,
                "base_sections": base_sections,
                "triggered_rules": hits,
                "insufficient_data_counts": dict(gap_counts),
                "confirmed_monthly_metrics": monthly_metrics,
                "official_macro_context": [
                    item for item in macro_context if item.get("kind") != "company_event"
                ],
                "company_risk_events": [
                    item for item in macro_context if item.get("kind") == "company_event"
                ],
                "previous_approved_style_examples": previous_style,
            },
            max_output_tokens=6000,
        )
    except Exception:
        return base_summary, base_sections, {}, "deterministic"
    if not result:
        return base_summary, base_sections, {}, "deterministic"
    payload = result.payload
    proposed = payload.get("sections")
    cited = payload.get("evidence_ids_by_section") or {}
    if not isinstance(proposed, dict) or any(
        not isinstance(proposed.get(key), str) for key in SECTION_ORDER
    ):
        return base_summary, base_sections, {}, "deterministic"
    for values in cited.values() if isinstance(cited, dict) else []:
        if not isinstance(values, list) or any(str(value) not in valid_evidence for value in values):
            return base_summary, base_sections, {}, "deterministic"
    summary = str(payload.get("summary") or base_summary)
    sections = {key: str(proposed[key])[:20_000] for key in SECTION_ORDER}
    return summary, sections, result.token_usage, result.model_name


def generate_monthly_report(db: Session, report_id: Any) -> MonthlyReport:
    report = db.get(MonthlyReport, report_id)
    if not report:
        raise ValueError("月报不存在。")
    company = db.get(Company, report.company_id)
    if not company:
        raise ValueError("公司不存在。")
    run = db.get(AnalysisRun, report.analysis_run_id) if report.analysis_run_id else None
    if not run or run.status != "completed":
        raise ValueError("请选择已完成的实时分析任务生成月报。")
    rows = _evaluations(db, run.id)
    if not rows:
        raise ValueError("分析任务尚无规则评估结果。")
    monthly_metrics = _monthly_metric_context(db, report)
    macro_context = _macro_context(db, report.period, company_id=report.company_id)
    as_of_date = _report_as_of_date(report.period)
    summary, sections, evidence = _deterministic_sections(
        company,
        report.period,
        rows,
        monthly_metrics,
        macro_context,
        as_of_date,
    )
    evidence.setdefault("finance", [])
    evidence["finance"].extend(
        str(metric["evidence_id"])
        for metric in monthly_metrics
        if metric.get("evidence_id")
    )
    summary, sections, token_usage, model_name = _llm_sections(
        db,
        report=report,
        company=company,
        rows=rows,
        base_summary=summary,
        base_sections=sections,
        evidence_by_section=evidence,
        monthly_metrics=monthly_metrics,
        macro_context=macro_context,
        as_of_date=as_of_date,
    )
    report.summary = summary
    report.sections = sections
    report.evidence_snapshot = {
        "by_section": evidence,
        "monthly_metrics": monthly_metrics,
        "macro_context": macro_context,
        "as_of_date": as_of_date.isoformat(),
        "period_status": "month_to_date" if as_of_date.day < calendar.monthrange(as_of_date.year, as_of_date.month)[1] else "closed_month",
        "rule_evaluation_ids": [str(evaluation.id) for evaluation, _ in rows],
    }
    report.status = "draft"
    report.model_name = model_name
    report.token_usage = token_usage
    report.generated_at = datetime.now(timezone.utc)
    db.add(report)
    db.flush()
    db.add(
        ReportVersion(
            report_id=report.id,
            revision=1,
            snapshot={
                "title": report.title,
                "summary": report.summary,
                "sections": report.sections,
                "evidence_snapshot": report.evidence_snapshot,
                "status": report.status,
            },
            edited_by=report.created_by,
        )
    )
    db.commit()
    db.refresh(report)
    return report


def append_report_revision(db: Session, report: MonthlyReport, actor: str) -> int:
    revision = (
        db.execute(
            select(func.max(ReportVersion.revision)).where(
                ReportVersion.report_id == report.id
            )
        ).scalar_one()
        or 0
    ) + 1
    db.add(
        ReportVersion(
            report_id=report.id,
            revision=revision,
            snapshot={
                "title": report.title,
                "summary": report.summary,
                "sections": report.sections,
                "evidence_snapshot": report.evidence_snapshot,
                "status": report.status,
            },
            edited_by=actor,
        )
    )
    return revision


def export_monthly_report_xlsx(
    db: Session, report: MonthlyReport
) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    company = db.get(Company, report.company_id)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "风险项目"
    sheet.sheet_view.showGridLines = False
    sheet.column_dimensions["A"].width = 22
    sheet.column_dimensions["B"].width = 110
    sheet.merge_cells("A1:B1")
    sheet["A1"] = report.title
    sheet["A1"].font = Font(name="Arial", size=16, bold=True, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor="17365D")
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 30
    sheet["A2"] = "企业名称"
    sheet["B2"] = company.name if company else ""
    sheet["A3"] = "监测期间"
    sheet["B3"] = report.period
    sheet["A4"] = "报告状态"
    sheet["B4"] = report.status
    sheet["A5"] = "监测截止日"
    sheet["B5"] = str((report.evidence_snapshot or {}).get("as_of_date") or "")
    thin = Side(style="thin", color="D9E2F3")
    header_fill = PatternFill("solid", fgColor="D9EAD3")
    row = 7
    for key in SECTION_ORDER:
        sheet.cell(row, 1, SECTION_LABELS[key])
        sheet.cell(row, 2, str((report.sections or {}).get(key) or ""))
        sheet.cell(row, 1).font = Font(name="Arial", bold=True, color="17365D")
        sheet.cell(row, 1).fill = header_fill
        sheet.cell(row, 1).alignment = Alignment(
            horizontal="center", vertical="top", wrap_text=True
        )
        sheet.cell(row, 2).alignment = Alignment(
            horizontal="left", vertical="top", wrap_text=True
        )
        sheet.cell(row, 1).border = Border(left=thin, right=thin, top=thin, bottom=thin)
        sheet.cell(row, 2).border = Border(left=thin, right=thin, top=thin, bottom=thin)
        content_length = len(str((report.sections or {}).get(key) or ""))
        sheet.row_dimensions[row].height = min(360, max(34, 20 + content_length / 3.5))
        row += 1
    sheet.freeze_panes = "A7"

    evaluation_sheet = workbook.create_sheet("规则评估")
    evaluation_sheet.append(
        ["规则编号", "风险板块", "监测要点", "评估结果", "重要性", "判断理由", "证据编号"]
    )
    for cell in evaluation_sheet[1]:
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for evaluation, rule in _evaluations(db, report.analysis_run_id):
        evaluation_sheet.append(
            [
                rule.code,
                rule.category,
                rule.monitoring_point,
                evaluation.status,
                evaluation.severity,
                evaluation.rationale,
                "、".join(
                    str(item.get("id") or "") for item in (evaluation.evidence or [])
                ),
            ]
        )
    widths = [12, 20, 58, 18, 14, 60, 38]
    for index, width in enumerate(widths, start=1):
        evaluation_sheet.column_dimensions[chr(64 + index)].width = width
    for row_cells in evaluation_sheet.iter_rows(min_row=2):
        for cell in row_cells:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    evaluation_sheet.freeze_panes = "A2"

    evidence_sheet = workbook.create_sheet("证据索引")
    evidence_sheet.append(["报告板块", "证据编号", "指标/说明", "期间", "数值", "来源位置"])
    for cell in evidence_sheet[1]:
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17365D")
    snapshot = report.evidence_snapshot or {}
    evidence_details = {
        str(item.get("evidence_id")): item
        for item in [
            *snapshot.get("monthly_metrics", []),
            *snapshot.get("macro_context", []),
        ]
        if item.get("evidence_id")
    }
    for section, evidence_ids in (snapshot.get("by_section") or {}).items():
        for evidence_id in evidence_ids:
            detail = evidence_details.get(str(evidence_id), {})
            value = detail.get("value_text", detail.get("value", detail.get("summary", "")))
            if detail.get("unit") and value not in (None, ""):
                value = f"{value}{detail['unit']}"
            source_locator = detail.get("source_locator", "")
            if not source_locator and detail.get("source_url"):
                source_locator = f"{detail.get('source_name') or ''} {detail['source_url']}".strip()
            evidence_sheet.append(
                [
                    SECTION_LABELS.get(section, section),
                    evidence_id,
                    detail.get("label", "规则评估证据，详见规则评估Sheet"),
                    detail.get("period", ""),
                    value,
                    source_locator,
                ]
            )
    for index, width in enumerate([20, 48, 40, 16, 22, 58], start=1):
        evidence_sheet.column_dimensions[chr(64 + index)].width = width
    evidence_sheet.freeze_panes = "A2"

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
