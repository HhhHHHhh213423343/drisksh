from __future__ import annotations

import hashlib
import io
import mimetypes
import re
import uuid
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from sqlalchemy import delete
from sqlalchemy.orm import Session

from app.config import get_settings
from app.models import ExtractedFact, SourceDocument
from app.services.llm import LLMService


ALLOWED_EXTENSIONS = {
    ".xlsx",
    ".xls",
    ".csv",
    ".pdf",
    ".docx",
    ".txt",
    ".png",
    ".jpg",
    ".jpeg",
}

METRIC_ALIASES: dict[str, tuple[str, ...]] = {
    "estimated_net_profit_impact": ("预计净利润影响", "政策变化净利润影响"),
    "overseas_asset_impairment_expected": ("境外资产减值预期", "境外资产预计减值比例"),
    "system_outage_days": ("核心系统中断天数", "业务系统中断天数"),
    "disaster_loss_to_equity": ("灾害损失占净资产", "不可抗力损失占净资产"),
    "revenue": ("营业收入", "主营业务收入", "营业总收入", "营收"),
    "net_profit": ("归母净利润", "净利润", "本期净利润"),
    "total_assets": ("资产总计", "总资产"),
    "total_liabilities": ("负债合计", "负债总计", "总负债"),
    "net_assets": ("所有者权益合计", "股东权益合计", "净资产"),
    "cash": ("货币资金", "现金及现金等价物"),
    "receivables": ("应收保理款", "应收账款", "应收款项"),
    "operating_cashflow": ("经营活动产生的现金流量净额", "经营现金净流量"),
    "top5_customer_ratio": ("前五大客户占比", "前五大交易对手方占比"),
    "top1_customer_ratio": ("第一大客户占比", "最大客户占比", "单一最大对手方占比"),
    "capital_operation_amount_to_equity": ("资本运作金额占净资产", "大额投融资占净资产"),
    "capital_operation_loss_to_equity": ("资本运作损失占净资产", "投资损失占净资产"),
    "core_revenue_yoy": ("核心业务收入同比", "核心业务营收同比"),
    "expense_growth_minus_revenue_growth": ("费用增速高于营收增速", "费用与营收增速差"),
    "related_party_transaction_ratio": ("关联交易占比",),
    "related_party_fund_occupation_to_equity": ("关联方资金占用占净资产", "资金占用余额占净资产"),
    "related_party_fund_occupation_days": ("关联方资金占用天数", "资金占用时间"),
    "valuation_yoy": ("估值同比", "估值变动率"),
    "revenue_yoy": ("营业收入同比", "主营业务收入同比", "营收同比"),
    "net_profit_yoy": ("净利润同比", "归母净利润同比"),
    "net_assets_yoy": ("净资产同比",),
    "core_business_metric_yoy": ("主要业务指标同比", "核心业务指标同比"),
    "roe_change_pp": ("ROE同比变动百分点", "净资产收益率变动百分点"),
    "cash_to_assets": ("货币资金占总资产", "现金占总资产"),
    "operating_cashflow_yoy": ("经营现金流同比", "经营活动现金流同比"),
    "overdue_90d_ratio_change_pp": ("90天以上逾期占比变动百分点", "逾期90天以上资产占比同比变动"),
    "cash_months_to_maturity": ("可动用资金覆盖到期债务月数", "现金覆盖到期债务月数"),
    "guarantees_to_equity": ("对外担保占净资产", "担保余额占净资产"),
    "accounting_change_profit_impact_abs": ("会计变更对净利润影响", "会计差错对净利润影响"),
    "key_role_vacancy_months": ("关键岗位缺位月数", "董事长总经理缺位月数"),
    "litigation_amount_to_equity": ("涉诉金额占净资产", "诉讼金额占净资产"),
    "pledged_share_ratio": ("股权质押比例", "质押股权占持股"),
    "frozen_share_ratio": ("股权冻结比例", "冻结股权占持股"),
    "employee_attrition_rate": ("员工离职率", "员工流失率"),
    "executive_turnover_rate": ("董监高离职率", "高管离职率"),
    "core_team_attrition_rate": ("核心团队流失率", "核心人员流失比例"),
    "top_shareholder_change_pp_abs": ("第一大股东持股变动百分点", "第一大股东持股比例变动"),
    "overdue_90d_ratio": ("逾期90天以上资产占比", "90天以上逾期占比"),
    "budget_completion": ("预算完成率",),
    "penalty_to_net_profit": ("处罚金额占净利润", "罚款占净利润"),
    "legal_guarantees_to_equity": ("法律担保余额占净资产", "连带保证责任占净资产"),
}

RATIO_METRICS = {
    code
    for code in METRIC_ALIASES
    if code.endswith("_ratio")
    or code.endswith("_yoy")
    or code
    in {
        "cash_to_assets",
        "budget_completion",
        "estimated_net_profit_impact",
        "overseas_asset_impairment_expected",
        "disaster_loss_to_equity",
        "capital_operation_amount_to_equity",
        "capital_operation_loss_to_equity",
        "expense_growth_minus_revenue_growth",
        "related_party_fund_occupation_to_equity",
        "guarantees_to_equity",
        "accounting_change_profit_impact_abs",
        "penalty_to_net_profit",
        "legal_guarantees_to_equity",
    }
}


@dataclass
class ParsedFact:
    metric_code: str
    label: str
    value_numeric: float | None
    value_text: str
    unit: str = ""
    currency: str = ""
    period: str = ""
    source_locator: str = ""
    confidence: float = 0.0
    raw_payload: dict[str, Any] = field(default_factory=dict)


@dataclass
class ParseResult:
    text: str
    facts: list[ParsedFact]
    page_count: int = 0
    metadata: dict[str, Any] = field(default_factory=dict)


def safe_filename(filename: str) -> str:
    basename = Path(filename or "upload").name
    cleaned = re.sub(r"[^\w\-.()\u4e00-\u9fff]+", "_", basename, flags=re.UNICODE)
    return cleaned[:180] or "upload"


def validate_extension(filename: str) -> str:
    extension = Path(filename).suffix.lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise ValueError(
            "仅支持 XLSX、XLS、CSV、PDF、DOCX、TXT、PNG、JPG 文件。"
        )
    return extension


def store_document_bytes(
    *, company_id: Any, filename: str, content: bytes
) -> tuple[str, str, str]:
    settings = get_settings()
    if len(content) > settings.max_upload_bytes:
        raise ValueError("文件超过 50MB 上传限制。")
    extension = validate_extension(filename)
    digest = hashlib.sha256(content).hexdigest()
    root = Path(settings.upload_root).expanduser().resolve()
    company_dir = (root / str(company_id)).resolve()
    if root not in company_dir.parents:
        raise ValueError("上传目录不合法。")
    company_dir.mkdir(parents=True, exist_ok=True)
    stored = company_dir / f"{uuid.uuid4().hex}{extension}"
    stored.write_bytes(content)
    return str(stored), digest, extension


def remove_stored_document(path: str) -> None:
    try:
        Path(path).unlink(missing_ok=True)
    except OSError:
        pass


def _normalize_label(value: Any) -> str:
    return re.sub(r"[\s（）()：:]", "", str(value or "")).lower()


def _metric_code(label: Any) -> str:
    normalized = _normalize_label(label)
    if not normalized:
        return ""
    matches: list[tuple[int, str]] = []
    for code, aliases in METRIC_ALIASES.items():
        for alias in aliases:
            normalized_alias = _normalize_label(alias)
            if normalized_alias and normalized_alias in normalized:
                matches.append((len(normalized_alias), code))
    return max(matches, default=(0, ""))[1]


def _number(value: Any, metric_code: str = "") -> tuple[float | None, str, str]:
    if isinstance(value, bool) or value is None:
        return None, "", ""
    if isinstance(value, (int, float)):
        return float(value), "", ""
    text = str(value).strip()
    if not text or text in {"-", "—", "/", "n.a.", "N/A"}:
        return None, text, ""
    negative = text.startswith("(") and text.endswith(")")
    match = re.search(r"[-+]?\d[\d,]*(?:\.\d+)?", text)
    if not match:
        return None, text, ""
    number = float(match.group(0).replace(",", ""))
    if negative:
        number = -abs(number)
    unit = ""
    if "%" in text:
        number /= 100
        unit = "%"
    elif "亿元" in text:
        number *= 100_000_000
        unit = "元"
    elif "万元" in text:
        number *= 10_000
        unit = "元"
    elif "千元" in text:
        number *= 1_000
        unit = "元"
    if metric_code in RATIO_METRICS and unit != "%" and abs(number) > 1:
        number /= 100
        unit = "%"
    if metric_code.endswith("_abs"):
        number = abs(number)
    return number, text, unit


def _period(value: Any) -> str:
    text = str(value or "").strip()
    match = re.search(r"(20\d{2})[年./-]?(0?[1-9]|1[0-2])?", text)
    if not match:
        return text[:32] if len(text) <= 32 else ""
    year, month = match.groups()
    return f"{year}-{int(month):02d}" if month else year


def _facts_from_rows(
    rows: list[list[Any]], *, source_name: str, base_confidence: float
) -> list[ParsedFact]:
    if not rows:
        return []
    facts: list[ParsedFact] = []
    header = rows[0]
    for row_index, row in enumerate(rows, start=1):
        if not row:
            continue
        for label_index, label in enumerate(row):
            code = _metric_code(label)
            if not code:
                continue
            for value_index in range(label_index + 1, len(row)):
                numeric, raw_text, unit = _number(row[value_index], code)
                if numeric is None:
                    continue
                period_value = header[value_index] if value_index < len(header) else ""
                facts.append(
                    ParsedFact(
                        metric_code=code,
                        label=str(label).strip(),
                        value_numeric=numeric,
                        value_text=raw_text or str(row[value_index]),
                        unit=unit,
                        currency="CNY" if unit == "元" else "",
                        period=_period(period_value),
                        source_locator=f"{source_name}!R{row_index}C{value_index + 1}",
                        confidence=base_confidence,
                        raw_payload={"row": row_index, "column": value_index + 1},
                    )
                )
            break

    # Also support tables where metrics are column headings and periods are rows.
    for column_index, label in enumerate(header):
        code = _metric_code(label)
        if not code:
            continue
        for row_index, row in enumerate(rows[1:], start=2):
            if column_index >= len(row):
                continue
            numeric, raw_text, unit = _number(row[column_index], code)
            if numeric is None:
                continue
            period_value = row[0] if row else ""
            facts.append(
                ParsedFact(
                    metric_code=code,
                    label=str(label).strip(),
                    value_numeric=numeric,
                    value_text=raw_text or str(row[column_index]),
                    unit=unit,
                    currency="CNY" if unit == "元" else "",
                    period=_period(period_value),
                    source_locator=f"{source_name}!R{row_index}C{column_index + 1}",
                    confidence=base_confidence,
                    raw_payload={"row": row_index, "column": column_index + 1},
                )
            )
    return _dedupe_facts(facts)


def _facts_from_text(text: str, *, source_name: str, confidence: float) -> list[ParsedFact]:
    facts: list[ParsedFact] = []
    for code, aliases in METRIC_ALIASES.items():
        alias_pattern = "|".join(re.escape(alias) for alias in aliases)
        pattern = re.compile(
            rf"(?P<label>{alias_pattern})[^\d负正()（）-]{{0,18}}(?P<value>[()（）+-]?\d[\d,]*(?:\.\d+)?\s*(?:亿元|万元|千元|元|%)?)"
        )
        for index, match in enumerate(pattern.finditer(text)):
            numeric, raw_text, unit = _number(match.group("value"), code)
            if numeric is None:
                continue
            context = text[max(0, match.start() - 30) : match.end() + 30]
            facts.append(
                ParsedFact(
                    metric_code=code,
                    label=match.group("label"),
                    value_numeric=numeric,
                    value_text=raw_text,
                    unit=unit,
                    currency="CNY" if unit == "元" else "",
                    period=_period(context),
                    source_locator=f"{source_name}#match-{index + 1}",
                    confidence=confidence,
                    raw_payload={"context": context},
                )
            )
    return _dedupe_facts(facts)


def _dedupe_facts(facts: list[ParsedFact]) -> list[ParsedFact]:
    selected: dict[tuple[str, str, float | None, str], ParsedFact] = {}
    for fact in facts:
        key = (fact.metric_code, fact.period, fact.value_numeric, fact.source_locator)
        current = selected.get(key)
        if current is None or current.confidence < fact.confidence:
            selected[key] = fact
    return list(selected.values())


def _parse_spreadsheet(path: Path, extension: str) -> ParseResult:
    import pandas as pd

    sheets: dict[str, Any]
    if extension == ".csv":
        try:
            frame = pd.read_csv(path, dtype=object)
        except UnicodeDecodeError:
            frame = pd.read_csv(path, dtype=object, encoding="gb18030")
        sheets = {"CSV": frame}
    else:
        sheets = pd.read_excel(path, sheet_name=None, dtype=object)
    text_parts: list[str] = []
    facts: list[ParsedFact] = []
    metadata: dict[str, Any] = {"sheets": []}
    for sheet_name, frame in list(sheets.items())[:50]:
        frame = frame.iloc[:5000, :200]
        frame = frame.where(frame.notna(), None)
        header = [str(column) for column in frame.columns]
        rows = [header, *frame.values.tolist()]
        text_rows = ["\t".join("" if value is None else str(value) for value in row) for row in rows]
        text_parts.append(f"[{sheet_name}]\n" + "\n".join(text_rows))
        facts.extend(
            _facts_from_rows(rows, source_name=str(sheet_name), base_confidence=0.96)
        )
        metadata["sheets"].append(
            {"name": str(sheet_name), "rows": len(frame), "columns": len(frame.columns)}
        )
    text = "\n\n".join(text_parts)[:120_000]
    facts.extend(_facts_from_text(text, source_name="workbook", confidence=0.82))
    return ParseResult(text=text, facts=_dedupe_facts(facts), metadata=metadata)


def _ocr_image(image: Any) -> str:
    try:
        import pytesseract

        return pytesseract.image_to_string(image, lang="chi_sim+eng")
    except Exception:
        return ""


def _parse_pdf(path: Path) -> ParseResult:
    from pypdf import PdfReader

    settings = get_settings()
    reader = PdfReader(str(path))
    page_count = len(reader.pages)
    if page_count > settings.max_document_pages:
        raise ValueError(f"PDF 超过 {settings.max_document_pages} 页限制。")
    pages: list[str] = []
    ocr_pages: list[int] = []
    for index, page in enumerate(reader.pages):
        text = (page.extract_text() or "").strip()
        if not text:
            try:
                import fitz
                from PIL import Image

                document = fitz.open(str(path))
                pixmap = document[index].get_pixmap(matrix=fitz.Matrix(1.8, 1.8), alpha=False)
                image = Image.open(io.BytesIO(pixmap.tobytes("png")))
                text = _ocr_image(image).strip()
                if text:
                    ocr_pages.append(index + 1)
                document.close()
            except Exception:
                text = ""
        pages.append(f"[第{index + 1}页]\n{text}")
    content = "\n\n".join(pages)[:120_000]
    confidence = 0.70 if ocr_pages else 0.84
    facts = _facts_from_text(content, source_name="pdf", confidence=confidence)
    return ParseResult(
        text=content,
        facts=facts,
        page_count=page_count,
        metadata={"ocr_pages": ocr_pages},
    )


def _parse_docx(path: Path) -> ParseResult:
    from docx import Document

    document = Document(str(path))
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    facts: list[ParsedFact] = []
    for table_index, table in enumerate(document.tables, start=1):
        rows = [[cell.text for cell in row.cells] for row in table.rows]
        parts.append(
            f"[表格{table_index}]\n"
            + "\n".join("\t".join(row) for row in rows)
        )
        facts.extend(
            _facts_from_rows(
                rows, source_name=f"table-{table_index}", base_confidence=0.92
            )
        )
    content = "\n".join(parts)[:120_000]
    facts.extend(_facts_from_text(content, source_name="docx", confidence=0.80))
    return ParseResult(text=content, facts=_dedupe_facts(facts))


def _parse_image(path: Path) -> ParseResult:
    from PIL import Image

    Image.MAX_IMAGE_PIXELS = 40_000_000
    image = Image.open(path)
    content = _ocr_image(image)[:120_000]
    if not content.strip():
        raise ValueError("图片 OCR 未提取到文字，请上传更清晰的文件。")
    return ParseResult(
        text=content,
        facts=_facts_from_text(content, source_name="image-ocr", confidence=0.68),
        page_count=1,
        metadata={"ocr": True, "width": image.width, "height": image.height},
    )


def parse_path(path: Path, extension: str) -> ParseResult:
    if extension in {".xlsx", ".docx"}:
        try:
            with zipfile.ZipFile(path) as archive:
                members = archive.infolist()
                total_size = sum(member.file_size for member in members)
                compressed_size = max(1, sum(member.compress_size for member in members))
                if len(members) > 10_000 or total_size > 250 * 1024 * 1024:
                    raise ValueError("Office 文件解压后体积过大，已拒绝解析。")
                if total_size / compressed_size > 150:
                    raise ValueError("Office 文件压缩比异常，已拒绝解析。")
        except zipfile.BadZipFile as exc:
            raise ValueError("Office 文件已损坏或格式不正确。") from exc
    if extension in {".xlsx", ".xls", ".csv"}:
        return _parse_spreadsheet(path, extension)
    if extension == ".pdf":
        return _parse_pdf(path)
    if extension == ".docx":
        return _parse_docx(path)
    if extension in {".png", ".jpg", ".jpeg"}:
        return _parse_image(path)
    if extension == ".txt":
        content = path.read_text(encoding="utf-8", errors="replace")[:120_000]
        return ParseResult(
            text=content,
            facts=_facts_from_text(content, source_name="text", confidence=0.78),
        )
    raise ValueError("不支持的文件类型。")


def parse_document(db: Session, document_id: Any) -> SourceDocument:
    document = db.get(SourceDocument, document_id)
    if not document:
        raise ValueError("材料不存在。")
    document.status = "parsing"
    document.parse_error = ""
    db.add(document)
    db.commit()
    try:
        result = parse_path(Path(document.storage_path), document.extension)
        llm_usage: dict[str, Any] = {}
        llm_mapping_error = ""
        if len(result.facts) < 5 and result.text.strip():
            service = LLMService()
            if service.is_configured:
                excerpt = (
                    result.text[:6000]
                    if len(result.text) <= 8000
                    else result.text[:5000] + "\n...[中间内容未发送]...\n" + result.text[-2000:]
                )
                try:
                    mapped = service.complete_json(
                        db,
                        purpose=f"document-fact-mapping:{document.sha256}",
                        system_prompt=(
                            "你是财务材料字段映射器。只能从输入摘录提取明确出现的数值，不能推算或补写。"
                            "返回JSON：{facts:[{metric_code,label,value_numeric,value_text,unit,currency,period,source_locator,confidence}]}。"
                            "metric_code必须从给定列表选择；不确定的字段不要返回。"
                        ),
                        user_payload={
                            "allowed_metric_codes": sorted(METRIC_ALIASES),
                            "filename": document.original_filename,
                            "text_excerpt": excerpt,
                        },
                        max_output_tokens=1800,
                    )
                except Exception as exc:
                    mapped = None
                    llm_mapping_error = str(exc)[:500]
                if mapped:
                    llm_usage = mapped.token_usage
                    for index, item in enumerate(mapped.payload.get("facts", [])):
                        if not isinstance(item, dict):
                            continue
                        code = str(item.get("metric_code") or "")
                        raw_value = (
                            item.get("value_numeric")
                            if item.get("value_numeric") is not None
                            else item.get("value_text")
                        )
                        value = _number(raw_value, code)[0]
                        if code not in METRIC_ALIASES or value is None:
                            continue
                        result.facts.append(
                            ParsedFact(
                                metric_code=code,
                                label=str(item.get("label") or code)[:255],
                                value_numeric=value,
                                value_text=str(item.get("value_text") or item.get("value_numeric") or ""),
                                unit=str(item.get("unit") or "")[:32],
                                currency=str(item.get("currency") or "")[:16],
                                period=str(item.get("period") or "")[:32],
                                source_locator=str(item.get("source_locator") or f"llm-excerpt-{index + 1}")[:255],
                                confidence=min(0.89, max(0.5, float(item.get("confidence") or 0.65))),
                                raw_payload={"mapping": "llm", "model": mapped.model_name},
                            )
                        )
                    result.facts = _dedupe_facts(result.facts)
        db.execute(delete(ExtractedFact).where(ExtractedFact.document_id == document.id))
        for fact in result.facts:
            status = "confirmed" if fact.confidence >= 0.92 else "pending"
            db.add(
                ExtractedFact(
                    document_id=document.id,
                    company_id=document.company_id,
                    metric_code=fact.metric_code,
                    label=fact.label,
                    value_numeric=fact.value_numeric,
                    value_text=fact.value_text,
                    unit=fact.unit,
                    currency=fact.currency,
                    period=fact.period,
                    source_locator=fact.source_locator,
                    confidence=fact.confidence,
                    status=status,
                    raw_payload=fact.raw_payload,
                    confirmed_by="automatic-parser" if status == "confirmed" else "",
                )
            )
        document.extracted_text = result.text
        document.page_count = result.page_count
        document.parse_metadata = {
            **result.metadata,
            "fact_count": len(result.facts),
            "auto_confirmed_count": sum(fact.confidence >= 0.92 for fact in result.facts),
            "llm_token_usage": llm_usage,
            "llm_mapping_error": llm_mapping_error,
        }
        document.status = "parsed"
    except Exception as exc:
        document.status = "failed"
        document.parse_error = str(exc)[:2000]
    db.add(document)
    db.commit()
    db.refresh(document)
    return document


def mime_type_for(filename: str, supplied: str = "") -> str:
    return supplied or mimetypes.guess_type(filename)[0] or "application/octet-stream"
